# =========================================
# التقارير: Excel جاهزة للتنزيل
# (الفواتير قابلة للطباعة/PDF من شاشة المحاسبة)
# =========================================
import io
from datetime import datetime

from fastapi import APIRouter, Depends
from fastapi.responses import Response
from sqlalchemy.orm import Session

from .. import models
from ..database import get_db

router = APIRouter(prefix="/reports", tags=["التقارير"])

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook(title: str, headers: list[str], rows: list[list]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.sheet_view.rightToLeft = True

    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for column_cells in ws.columns:
        width = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
        ws.column_dimensions[column_cells[0].column_letter].width = min(width + 4, 45)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_response(name: str, content: bytes) -> Response:
    stamp = datetime.utcnow().strftime("%Y%m%d")
    return Response(content=content, media_type=XLSX_TYPE,
                    headers={"Content-Disposition": f'attachment; filename="{name}-{stamp}.xlsx"'})


@router.get("/inventory.xlsx")
def inventory_report(db: Session = Depends(get_db)):
    """تقرير المخزون الكامل"""
    rows = [
        [p.id, p.name, p.sku,
         "مادة خام" if p.product_type == "raw" else "منتج نهائي",
         p.category, p.quantity, p.unit, p.reorder_point,
         p.unit_cost, p.unit_price, round(p.quantity * p.unit_cost, 2),
         "أعد الطلب" if p.quantity <= p.reorder_point else "متوفر"]
        for p in db.query(models.Product).all()
    ]
    content = _workbook("المخزون",
                        ["#", "الاسم", "SKU", "النوع", "التصنيف", "الكمية", "الوحدة",
                         "حد الطلب", "التكلفة", "سعر البيع", "قيمة المخزون", "الحالة"], rows)
    return _xlsx_response("inventory", content)


@router.get("/sales.xlsx")
def sales_report(db: Session = Depends(get_db)):
    """تقرير المبيعات الكامل"""
    products = {p.id: p.name for p in db.query(models.Product).all()}
    status_ar = {"confirmed": "مؤكد", "pending_production": "بانتظار الإنتاج",
                 "delivered": "مُسلَّم", "cancelled": "ملغى"}
    rows = [
        [o.id, products.get(o.product_id, o.product_id), o.customer_name,
         o.quantity, o.unit_price, round(o.quantity * o.unit_price, 2),
         status_ar.get(o.status, o.status), o.ordered_at.strftime("%Y-%m-%d")]
        for o in db.query(models.SalesOrder).order_by(models.SalesOrder.id.desc()).all()
    ]
    content = _workbook("المبيعات",
                        ["#", "المنتج", "العميل", "الكمية", "سعر الوحدة",
                         "الإجمالي", "الحالة", "التاريخ"], rows)
    return _xlsx_response("sales", content)


@router.get("/journal.xlsx")
def journal_report(db: Session = Depends(get_db)):
    """تقرير القيود المحاسبية"""
    rows = [
        [e.id, e.description, e.debit_account, e.credit_account,
         e.amount, e.reference, e.created_at.strftime("%Y-%m-%d %H:%M")]
        for e in db.query(models.JournalEntry).order_by(models.JournalEntry.id.desc()).all()
    ]
    content = _workbook("القيود المحاسبية",
                        ["#", "البيان", "مدين", "دائن", "المبلغ", "المرجع", "التاريخ"], rows)
    return _xlsx_response("journal", content)


@router.get("/attendance.xlsx")
def attendance_report(db: Session = Depends(get_db)):
    """تقرير الحضور والانصراف"""
    employees = {e.id: e for e in db.query(models.Employee).all()}
    rows = []
    for r in db.query(models.AttendanceRecord).order_by(models.AttendanceRecord.check_in.desc()).all():
        emp = employees.get(r.employee_id)
        hours = round((r.check_out - r.check_in).total_seconds() / 3600, 2) if r.check_out else None
        rows.append([
            emp.name if emp else r.employee_id,
            emp.department if emp else "",
            r.check_in.strftime("%Y-%m-%d %H:%M"),
            r.check_out.strftime("%Y-%m-%d %H:%M") if r.check_out else "لم ينصرف",
            hours if hours is not None else "",
            "جهاز بصمة" if r.source == "device" else "يدوي",
        ])
    content = _workbook("الحضور",
                        ["الموظف", "القسم", "الحضور", "الانصراف", "الساعات", "المصدر"], rows)
    return _xlsx_response("attendance", content)
