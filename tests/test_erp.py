# =========================================
# اختبارات نظام Smart Factory ERP
# =========================================


# ---------- الأساسيات واللوحة ----------
class TestBasics:
    def test_dashboard_html_at_root(self, client):
        response = client.get("/")
        assert response.status_code == 200
        assert "Smart Factory ERP" in response.text

    def test_api_info(self, client):
        assert client.get("/api/info").status_code == 200

    def test_dashboard_kpis(self, client):
        data = client.get("/dashboard").json()
        assert data["inventory"]["products_count"] >= 4
        assert data["machines"]["total"] >= 3

    def test_sales_timeseries(self, client):
        data = client.get("/dashboard/sales-timeseries?days=30").json()
        assert data["series"], "يجب أن توجد مبيعات يومية في بيانات seed"


# ---------- المصادقة والصلاحيات ----------
class TestAuth:
    def test_bad_login_rejected(self, client):
        response = client.post("/auth/login", json={"username": "admin", "password": "wrong"})
        assert response.status_code == 401

    def test_write_without_token_rejected(self, client):
        response = client.post("/inventory/products", json={"name": "x", "sku": "NOAUTH"})
        assert response.status_code == 401

    def test_garbage_token_rejected(self, client):
        response = client.post("/employees", json={"name": "x"},
                               headers={"Authorization": "Bearer abc.def"})
        assert response.status_code == 401

    def test_viewer_cannot_write(self, client, admin_headers):
        client.post("/auth/register", headers=admin_headers, json={
            "username": "viewer_t", "password": "secret1", "role": "viewer"})
        login = client.post("/auth/login", json={"username": "viewer_t", "password": "secret1"})
        viewer_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        assert client.get("/inventory/products", headers=viewer_headers).status_code == 200
        assert client.post("/employees", headers=viewer_headers,
                           json={"name": "x"}).status_code == 403

    def test_admin_lists_users(self, client, admin_headers):
        assert client.get("/auth/users", headers=admin_headers).status_code == 200


# ---------- المخزون ----------
class TestInventory:
    def test_create_and_duplicate_sku(self, client, admin_headers):
        product = {"name": "منتج اختبار", "sku": "TST-1", "quantity": 50, "unit_price": 9}
        assert client.post("/inventory/products", headers=admin_headers, json=product).status_code == 200
        assert client.post("/inventory/products", headers=admin_headers, json=product).status_code == 400

    def test_over_withdraw_rejected(self, client, admin_headers):
        response = client.post("/inventory/products/1/adjust", headers=admin_headers,
                               json={"change": -999999})
        assert response.status_code == 400

    def test_alerts(self, client):
        data = client.get("/inventory/alerts").json()
        assert "alerts" in data


# ---------- تعديل وحذف (CRUD) ----------
class TestCrud:
    def test_update_product(self, client, admin_headers):
        created = client.post("/inventory/products", headers=admin_headers, json={
            "name": "قبل التعديل", "sku": "UPD-1", "quantity": 10}).json()
        updated = client.put(f"/inventory/products/{created['id']}", headers=admin_headers,
                             json={"name": "بعد التعديل", "unit_price": 99}).json()
        assert updated["name"] == "بعد التعديل" and updated["unit_price"] == 99
        assert updated["quantity"] == 10  # الكمية لا تتغير من التعديل

    def test_delete_product_without_moves(self, client, admin_headers):
        created = client.post("/inventory/products", headers=admin_headers, json={
            "name": "للحذف", "sku": "DEL-1"}).json()
        assert client.delete(f"/inventory/products/{created['id']}",
                             headers=admin_headers).status_code == 200
        assert client.get(f"/inventory/products/{created['id']}").status_code == 404

    def test_delete_product_with_sales_blocked(self, client, admin_headers):
        # المنتج 1 له مبيعات في seed — يجب رفض حذفه
        assert client.delete("/inventory/products/1", headers=admin_headers).status_code == 400

    def test_update_machine(self, client, admin_headers):
        created = client.post("/machines", headers=admin_headers, json={"name": "آلة قبل"}).json()
        updated = client.put(f"/machines/{created['id']}", headers=admin_headers,
                             json={"name": "آلة بعد", "status": "stopped"}).json()
        assert updated["name"] == "آلة بعد" and updated["status"] == "stopped"

    def test_delete_machine_without_orders(self, client, admin_headers):
        created = client.post("/machines", headers=admin_headers, json={"name": "آلة للحذف"}).json()
        assert client.delete(f"/machines/{created['id']}", headers=admin_headers).status_code == 200

    def test_delete_machine_with_orders_blocked(self, client, admin_headers):
        # الآلة 1 لها أوامر إنتاج في seed
        assert client.delete("/machines/1", headers=admin_headers).status_code == 400

    def test_update_employee(self, client, admin_headers):
        created = client.post("/employees", headers=admin_headers,
                              json={"name": "موظف", "salary": 1000}).json()
        updated = client.put(f"/employees/{created['id']}", headers=admin_headers,
                             json={"salary": 2000}).json()
        assert updated["salary"] == 2000 and updated["name"] == "موظف"

    def test_crud_requires_writer(self, client):
        assert client.put("/inventory/products/1", json={"name": "x"}).status_code == 401
        assert client.delete("/machines/1").status_code == 401


# ---------- الآلات والصيانة التنبؤية ----------
class TestMachines:
    def test_risk_overview(self, client):
        machines = client.get("/machines/risk-overview").json()["machines"]
        assert machines and all("risk_level" in m for m in machines)

    def test_stressed_machine_high_risk(self, client):
        # الآلة 2 في بيانات seed مجهدة عمدًا (حرارة واهتزاز مرتفعان)
        data = client.get("/machines/2/predict-failure").json()
        assert data["failure_probability"] > 0.5

    def test_healthy_machine_low_risk(self, client):
        data = client.get("/machines/1/predict-failure").json()
        assert data["failure_probability"] < 0.4

    def test_missing_machine_404(self, client):
        assert client.get("/machines/999/predict-failure").status_code == 404


# ---------- الإنتاج والمبيعات ----------
class TestProductionAndSales:
    def test_production_cycle_adds_to_stock(self, client, admin_headers):
        before = client.get("/inventory/products/1").json()["quantity"]
        order = client.post("/production/orders", headers=admin_headers,
                            json={"product_id": 1, "machine_id": 1, "planned_quantity": 100}).json()
        client.post(f"/production/orders/{order['id']}/complete", headers=admin_headers,
                    json={"produced_quantity": 90, "defective_quantity": 5})
        after = client.get("/inventory/products/1").json()["quantity"]
        assert after == before + 85  # الصافي = المنتَج - المعيب

    def test_quality_anomalies(self, client):
        data = client.get("/production/quality/anomalies").json()
        assert data["status"] == "ok"
        assert data["anomalies_found"] >= 2  # أمران شاذان مزروعان في seed

    def test_oversell_triggers_auto_production(self, client, admin_headers):
        """بيع فوق المتاح ← طلب بانتظار الإنتاج + أمر إنتاج تلقائي للعجز"""
        stock = client.get("/inventory/products/1").json()["quantity"]
        order = client.post("/sales/orders", headers=admin_headers,
                            json={"product_id": 1, "quantity": stock + 500}).json()
        assert order["status"] == "pending_production"
        auto = [o for o in client.get("/production/orders").json()
                if o.get("sales_order_id") == order["id"]]
        assert len(auto) == 1 and auto[0]["source"] == "auto"
        assert auto[0]["planned_quantity"] == 500

    def test_sale_deducts_stock(self, client, admin_headers):
        before = client.get("/inventory/products/1").json()["quantity"]
        client.post("/sales/orders", headers=admin_headers, json={"product_id": 1, "quantity": 5})
        after = client.get("/inventory/products/1").json()["quantity"]
        assert after == before - 5

    def test_demand_forecast(self, client):
        data = client.get("/sales/forecast/1?days_ahead=30").json()
        assert data["status"] == "ok"
        assert data["expected_total_demand"] > 0


# ---------- السلسلة التلقائية والمحاسبة والمشتريات ----------
class TestWorkflow:
    def test_departments_permissions(self, client, admin_headers):
        """مستخدم قسم المبيعات يبيع لكنه لا يضيف منتجات (قسم المخازن)"""
        login = client.post("/auth/login", json={"username": "sales", "password": "sales123"})
        assert login.status_code == 200
        sales_headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        ok = client.post("/sales/orders", headers=sales_headers,
                         json={"product_id": 1, "quantity": 1})
        assert ok.status_code == 200
        denied = client.post("/inventory/products", headers=sales_headers,
                             json={"name": "x", "sku": "DENY-1"})
        assert denied.status_code == 403

    def test_deliver_issues_invoice_and_journal(self, client, admin_headers):
        order = client.post("/sales/orders", headers=admin_headers,
                            json={"product_id": 1, "quantity": 3, "customer_name": "عميل الفاتورة"}).json()
        assert order["status"] == "confirmed"
        result = client.post(f"/sales/orders/{order['id']}/deliver", headers=admin_headers).json()
        assert result["invoice_number"].startswith("INV-")
        invoices = client.get("/accounting/invoices").json()
        assert any(i["number"] == result["invoice_number"] for i in invoices)
        journal = client.get("/accounting/journal").json()
        refs = [e for e in journal if e["reference"] == result["invoice_number"]]
        assert any(e["entry_type"] == "sale" for e in refs)
        assert any(e["entry_type"] == "cogs" for e in refs)
        # الفاتورة قابلة للطباعة
        inv_id = next(i["id"] for i in invoices if i["number"] == result["invoice_number"])
        printable = client.get(f"/accounting/invoices/{inv_id}/print")
        assert printable.status_code == 200 and "فاتورة" in printable.text

    def test_full_auto_chain(self, client, admin_headers):
        """السلسلة الكاملة: بيع بعجز ← إنتاج تلقائي ← نقص مواد خام ← طلب شراء تلقائي
        ← اعتماد واستلام ← إكمال الإنتاج ← تسليم وفاتورة"""
        # منتج جديد بمخزون صفر + مادة خام شحيحة + BOM
        raw = client.post("/inventory/products", headers=admin_headers, json={
            "name": "مادة خام للاختبار", "sku": "RAW-T1", "product_type": "raw",
            "quantity": 5, "unit_cost": 2}).json()
        finished = client.post("/inventory/products", headers=admin_headers, json={
            "name": "منتج السلسلة", "sku": "CHAIN-1", "quantity": 0, "unit_price": 50}).json()
        client.post(f"/inventory/products/{finished['id']}/bom", headers=admin_headers,
                    json={"component_id": raw["id"], "quantity_per_unit": 2})

        # 1) بيع 10 والمخزون صفر ← pending_production
        sale = client.post("/sales/orders", headers=admin_headers,
                           json={"product_id": finished["id"], "quantity": 10}).json()
        assert sale["status"] == "pending_production"

        # 2) أمر إنتاج تلقائي أُنشئ
        production = next(o for o in client.get("/production/orders").json()
                          if o.get("sales_order_id") == sale["id"])

        # 3) المواد الخام لا تكفي (نحتاج 20 والمتاح 5) ← طلب شراء تلقائي
        auto_po = next(p for p in client.get("/purchases/orders").json()
                       if p["product_id"] == raw["id"] and p["source"] == "auto")
        assert auto_po["status"] == "requested" and auto_po["supplier_id"] is None

        # 4) لا يمكن إكمال الإنتاج قبل توفر المواد
        blocked = client.post(f"/production/orders/{production['id']}/complete",
                              headers=admin_headers,
                              json={"produced_quantity": 10, "defective_quantity": 0})
        assert blocked.status_code == 400

        # 5) قسم المشتريات يعتمد الطلب ويستلمه
        supplier_id = client.get("/purchases/suppliers").json()[0]["id"]
        approved = client.post(f"/purchases/orders/{auto_po['id']}/approve", headers=admin_headers,
                               json={"supplier_id": supplier_id}).json()
        assert approved["status"] == "ordered"
        received = client.post(f"/purchases/orders/{auto_po['id']}/receive", headers=admin_headers).json()
        assert received["status"] == "received"

        # 6) الآن يكتمل الإنتاج: تُستهلك المواد ويدخل المنتج للمخزون
        raw_before = client.get(f"/inventory/products/{raw['id']}").json()["quantity"]
        done = client.post(f"/production/orders/{production['id']}/complete", headers=admin_headers,
                           json={"produced_quantity": 10, "defective_quantity": 0})
        assert done.status_code == 200
        raw_after = client.get(f"/inventory/products/{raw['id']}").json()["quantity"]
        assert raw_after == raw_before - 20  # استهلاك BOM: 2 لكل وحدة

        # 7) التسليم: فاتورة + قيود
        delivered = client.post(f"/sales/orders/{sale['id']}/deliver", headers=admin_headers)
        assert delivered.status_code == 200
        assert delivered.json()["invoice_number"].startswith("INV-")

    def test_purchase_receipt_creates_journal_entry(self, client, admin_headers):
        supplier_id = client.get("/purchases/suppliers").json()[0]["id"]
        po = client.post("/purchases/orders", headers=admin_headers, json={
            "supplier_id": supplier_id, "product_id": 1, "quantity": 10, "unit_cost": 7}).json()
        client.post(f"/purchases/orders/{po['id']}/receive", headers=admin_headers)
        journal = client.get("/accounting/journal").json()
        assert any(e["reference"] == f"PO-{po['id']}" and e["entry_type"] == "purchase"
                   for e in journal)

    def test_financial_summary(self, client):
        summary = client.get("/accounting/summary").json()
        assert summary["revenue"] > 0
        assert "gross_profit" in summary

    def test_supplier_crud(self, client, admin_headers):
        created = client.post("/purchases/suppliers", headers=admin_headers,
                              json={"name": "مورد اختبار"}).json()
        updated = client.put(f"/purchases/suppliers/{created['id']}", headers=admin_headers,
                             json={"phone": "0555"}).json()
        assert updated["phone"] == "0555"
        assert client.delete(f"/purchases/suppliers/{created['id']}",
                             headers=admin_headers).status_code == 200


# ---------- الحضور والتكامل والتقارير والباركود ----------
class TestPhase2:
    def test_manual_attendance_cycle(self, client, admin_headers):
        r = client.post("/hr/attendance/1/in", headers=admin_headers)
        assert r.status_code == 200
        # حضور مكرر مرفوض
        assert client.post("/hr/attendance/1/in", headers=admin_headers).status_code == 400
        day = client.get("/hr/attendance").json()
        record = next(x for x in day["records"] if x["employee_id"] == 1)
        assert record["status"] == "✅ حاضر"
        assert client.post("/hr/attendance/1/out", headers=admin_headers).status_code == 200
        day = client.get("/hr/attendance").json()
        record = next(x for x in day["records"] if x["employee_id"] == 1)
        assert record["status"] == "🏠 انصرف" and record["hours"] is not None

    def test_api_key_and_device_attendance(self, client, admin_headers):
        # جهاز بلا مفتاح مرفوض
        assert client.post("/integrations/attendance",
                           json={"badge_code": "1002", "direction": "in"}).status_code == 401
        # إنشاء مفتاح (admin فقط)
        created = client.post("/integrations/keys", headers=admin_headers,
                              json={"name": "جهاز بصمة تجريبي"}).json()
        api_key = created["api_key"]
        assert api_key.startswith("erp_")
        # الجهاز يسجل حضورًا بالبطاقة
        r = client.post("/integrations/attendance", headers={"X-API-Key": api_key},
                        json={"badge_code": "1002", "direction": "in"})
        assert r.status_code == 200
        day = client.get("/hr/attendance").json()
        record = next(x for x in day["records"] if x["badge_code"] == "1002")
        assert record["source"] == "device"
        # بطاقة غير معروفة
        assert client.post("/integrations/attendance", headers={"X-API-Key": api_key},
                           json={"badge_code": "9999", "direction": "in"}).status_code == 404
        # إيقاف المفتاح يمنع الجهاز
        key_id = client.get("/integrations/keys", headers=admin_headers).json()[-1]["id"]
        client.delete(f"/integrations/keys/{key_id}", headers=admin_headers)
        assert client.post("/integrations/attendance", headers={"X-API-Key": api_key},
                           json={"badge_code": "1003", "direction": "in"}).status_code == 401

    def test_product_qrcode(self, client):
        r = client.get("/integrations/qrcode/product/1")
        assert r.status_code == 200
        assert r.headers["content-type"] == "image/png"
        assert r.content[:8] == b"\x89PNG\r\n\x1a\n"

    def test_excel_reports(self, client):
        from io import BytesIO
        from openpyxl import load_workbook
        for path in ["/reports/inventory.xlsx", "/reports/sales.xlsx",
                     "/reports/journal.xlsx", "/reports/attendance.xlsx"]:
            r = client.get(path)
            assert r.status_code == 200, path
            wb = load_workbook(BytesIO(r.content))
            assert wb.active.max_row >= 1, path


# ---------- التحليلات الذكية ----------
class TestAnalytics:
    def test_analytics_endpoint_structure(self, client):
        data = client.get("/ai/analytics").json()
        for key in ("raw_materials", "profitability", "delays", "machines", "employees"):
            assert key in data, key

    def test_raw_material_runway(self, client):
        materials = client.get("/ai/analytics").json()["raw_materials"]
        assert materials, "يجب أن توجد مواد خام في seed"
        assert all("days_until_stockout" in m and "daily_consumption" in m for m in materials)

    def test_profitability_sorted_desc(self, client):
        rows = client.get("/ai/analytics").json()["profitability"]
        profits = [r["profit"] for r in rows]
        assert profits == sorted(profits, reverse=True)
        assert rows[0]["revenue"] > 0

    def test_machine_performance_scores(self, client):
        machines = client.get("/ai/analytics").json()["machines"]
        assert machines
        assert all(0 <= m["performance_score"] <= 100 for m in machines)

    def test_delays_analysis(self, client, admin_headers):
        delays = client.get("/ai/analytics").json()["delays"]
        assert "late_orders" in delays and "message" in delays

    def test_raw_material_shortage_in_insights(self, client, admin_headers):
        """مادة خام يحتاجها أمر مفتوح أكثر من المتاح ← تظهر في الرؤى"""
        raw = client.post("/inventory/products", headers=admin_headers, json={
            "name": "خام الرؤى", "sku": "RAW-INS", "product_type": "raw",
            "quantity": 1, "unit_cost": 1}).json()
        fin = client.post("/inventory/products", headers=admin_headers, json={
            "name": "منتج الرؤى", "sku": "FIN-INS", "quantity": 0}).json()
        client.post(f"/inventory/products/{fin['id']}/bom", headers=admin_headers,
                    json={"component_id": raw["id"], "quantity_per_unit": 5})
        client.post("/production/orders", headers=admin_headers,
                    json={"product_id": fin["id"], "planned_quantity": 100})
        insights = client.get("/ai/insights").json()["insights"]
        assert any(i["category"] == "مواد خام" and "خام الرؤى" in i["title"] for i in insights)


# ---------- الذكاء الاصطناعي ----------
class TestAI:
    def test_assistant_fallback_engine(self, client):
        data = client.post("/ai/assistant", json={"question": "كيف حال المخزون؟"}).json()
        assert data["engine"] == "rules"
        assert data["answer"]

    def test_assistant_topics(self, client):
        for question in ["هل توجد أعطال في الآلات؟", "كم إيرادات المبيعات؟", "ما وضع الإنتاج؟"]:
            data = client.post("/ai/assistant", json={"question": question}).json()
            assert data["answer"]

    def test_insights_sorted_by_priority(self, client):
        data = client.get("/ai/insights").json()
        assert data["insights_count"] >= 1
        priorities = [i["priority"] for i in data["insights"]]
        assert priorities == sorted(priorities)

    def test_insights_include_maintenance(self, client):
        data = client.get("/ai/insights").json()
        assert any(i["category"] == "صيانة" for i in data["insights"])


# ---------- إنترنت الأشياء ----------
class TestIoT:
    def test_parse_topic(self):
        from erp.iot.mqtt_ingest import parse_topic_machine_id
        assert parse_topic_machine_id("factory/machines/3/sensors") == 3
        assert parse_topic_machine_id("factory/other/3/sensors") is None
        assert parse_topic_machine_id("factory/machines/abc/sensors") is None

    def test_ingest_valid_reading(self, client):
        from erp.iot.mqtt_ingest import ingest_reading
        result = ingest_reading(
            "factory/machines/1/sensors",
            b'{"temperature": 71, "vibration": 2.5, "pressure": 5.1, "running_hours": 210}',
        )
        assert result["status"] == "ok"

    def test_ingest_rejects_bad_payload(self):
        from erp.iot.mqtt_ingest import ingest_reading
        assert ingest_reading("factory/machines/1/sensors", b"not json")["status"] == "error"
        assert ingest_reading("factory/machines/1/sensors", b'{"temperature": 70}')["status"] == "error"

    def test_ingest_rejects_unknown_machine(self):
        from erp.iot.mqtt_ingest import ingest_reading
        result = ingest_reading(
            "factory/machines/999/sensors",
            b'{"temperature": 70, "vibration": 2, "pressure": 5, "running_hours": 10}',
        )
        assert result["status"] == "error"

    def test_listener_off_without_env(self):
        from erp.iot.mqtt_ingest import start_mqtt_listener
        assert start_mqtt_listener() is False
